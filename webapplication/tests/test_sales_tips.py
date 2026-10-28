"""Sales, tips/service charges, credit collections, reconciliation, report totals."""
from decimal import Decimal
import pytest
from services.accounting_service import DuplicateWarning
from services.reconciliation_service import ReconciliationService


def day(svc, **kw):
    base = dict(business_date="01-10-2026", gross_sales=1000, discounts=100, refunds=0, cash=550, upi=300,
                card=100, tips=50, service_charges=0)
    base.update(kw)
    return svc.add_sales(**base)


def test_sales_reconcile_with_tips_included_in_collections(svc):
    # net 900 + tips 50 = 950 collected
    day(svc)
    r = ReconciliationService(svc).sales_rows()[0]
    assert r["expected_receipts"] == Decimal("950.00") and r["difference"] == 0 and r["status"] == "OK"


def test_reconciliation_flags_difference(svc):
    day(svc, cash=400)
    r = ReconciliationService(svc).sales_rows()[0]
    assert r["difference"] == Decimal("-150.00") and "Collections differ" in r["status"]


def test_tips_and_service_charges_do_not_inflate_revenue(svc):
    day(svc, tips=50, service_charges=40, cash=490)
    s = svc.dashboard_summary()
    assert s["net_sales"] == Decimal("900.00") and s["cafe_revenue"] == Decimal("900.00")
    assert s["tips"] == Decimal("50.00") and s["service_charges"] == Decimal("40.00")
    svc.config["tips_excluded_from_revenue"] = False          # accountant policy change
    assert svc.dashboard_summary()["cafe_revenue"] == Decimal("950.00")


def test_tips_ledger_entries_separate_and_linked(svc):
    sid = day(svc, tips=50, service_charges=40, cash=490)
    led = svc.db.read("tips_ledger")
    assert sorted(led["entry_type"]) == ["service_charge", "tip"]
    assert set(led["sales_id"]) == {sid}


def test_credit_collections_are_not_new_revenue(svc):
    day(svc, business_date="02-10-2026", gross_sales=500, discounts=0, cash=300, upi=0, card=0, tips=0,
        credit_sales=200, credit_collections=150)
    # receipts 300 = net 500 - credit sales 200 + 0 ... plus collections 150 => expected 450; entered 300
    s = svc.dashboard_summary()
    assert s["net_sales"] == Decimal("500.00") and s["cafe_revenue"] == Decimal("500.00")
    day(svc, business_date="03-10-2026", gross_sales=0, discounts=0, cash=150, upi=0, card=0, tips=0,
        credit_collections=150)
    assert svc.dashboard_summary()["cafe_revenue"] == Decimal("500.00")     # unchanged by collections
    assert ReconciliationService(svc).sales_rows("03-10-2026", "03-10-2026")[0]["difference"] == 0


def test_supplier_payments_not_counted_as_expenses_twice(svc, rep, vendor):
    b = svc.add_bill(vendor, "E1", "01-10-2026", "31-10-2026", "Milk and Dairy", 1000)
    svc.record_payment(vendor, "05-10-2026", 1000, "Cash", [(b, 1000)])
    assert rep.expense_breakdown().iloc[0]["invoice_total"] == Decimal("1000.00")   # bill counted once
    cf = rep.cash_flow()
    assert cf["amount"].sum() == Decimal("-1000.00")


def test_duplicate_date_shift_warns_but_allows_confirmed_and_multishift(svc):
    day(svc)
    with pytest.raises(DuplicateWarning):
        day(svc)
    day(svc, allow_duplicate=True)
    svc.config["allow_multiple_shifts"] = True
    day(svc, business_date="05-10-2026", shift="Morning")
    day(svc, business_date="05-10-2026", shift="Evening")          # legitimate, no warning
    with pytest.raises(DuplicateWarning):
        day(svc, business_date="05-10-2026", shift="Evening")


def test_edit_sales_syncs_tip_ledger_and_is_audited(svc):
    sid = day(svc, tips=50, cash=500)
    svc.update_sales(sid, tips_collected="80", reason="recount")
    led = svc.db.read("tips_ledger")
    assert led.iloc[0]["collected"] == "80.00" and led.iloc[0]["retained_or_payable"] == "80.00"
    assert ReconciliationService(svc).sales_rows()[0]["tip_ledger_diff"] == 0
    with pytest.raises(ValueError):
        svc.update_sales(sid, status="Voided")
    a = svc.db.read("audit_log")
    last = a[a["action"] == "update"].iloc[-1]
    assert "50.00" in last["before"] and "80.00" in last["after"]


def test_tip_distribution_rules_and_void_cascade(svc):
    sid = day(svc, tips=50, cash=500)
    eid = svc.db.read("tips_ledger").iloc[0]["entry_id"]
    with pytest.raises(ValueError):
        svc.update_tip(eid, 20, 20)                                    # does not sum to collected
    svc.update_tip(eid, 20, 30, "paid staff", "02-10-2026")
    with pytest.raises(ValueError, match="distributed"):
        svc.update_sales(sid, tips_collected="10")
    with pytest.raises(ValueError, match="distributed"):
        svc.void_sales(sid, "mistake")
    svc.update_tip(eid, 0, 50)
    svc.void_sales(sid, "mistake")
    assert svc.db.read("customer_sales").iloc[0]["status"] == "Voided"
    assert svc.db.read("tips_ledger").iloc[0]["status"] == "Voided"


def test_unrecorded_actual_cash_is_not_fabricated(svc):
    day(svc)
    row = svc.db.read("customer_sales").iloc[0]
    assert row["actual_cash"] == "" and row["cash_difference"] == ""
    day(svc, business_date="09-10-2026", actual_cash=545)
    assert svc.db.read("customer_sales").iloc[1]["cash_difference"] == "-5.00"


def test_sales_adjustment_never_overwrites_original(svc):
    sid = day(svc)
    aid = svc.add_sales_adjustment(sid, "REFUND", 100, "05-10-2026", "late refund")
    assert svc.db.read("customer_sales").iloc[0]["refunds"] == "0.00"
    assert svc.dashboard_summary()["net_sales"] == Decimal("800.00")
    svc.void_sales_adjustment(aid, "entered twice")
    assert svc.dashboard_summary()["net_sales"] == Decimal("900.00")


def test_reports_match_ledgers(svc, rep, vendor):
    b1 = svc.add_bill(vendor, "R1", "01-10-2026", "31-10-2026", "Milk and Dairy", 1200)
    b2 = svc.add_bill(vendor, "R2", "15-11-2026", "30-11-2026", "Other", 800)
    svc.record_payment(vendor, "05-10-2026", 500, "UPI", [(b1, 500)])
    day(svc)
    day(svc, business_date="02-10-2026", gross_sales=2000, discounts=0, cash=2000, upi=0, card=0, tips=0)
    assert rep.expense_breakdown()["invoice_total"].sum() == Decimal("2000.00")
    assert rep.expense_periodic("M")["invoice_total"].sum() == Decimal("2000.00")
    assert rep.expense_periodic("Q")["invoice_total"].sum() == Decimal("2000.00")
    assert rep.expense_by_vendor()["invoice_total"].sum() == Decimal("2000.00")
    assert rep.vendor_outstanding_summary("30-11-2026")["total_outstanding"].sum() == Decimal("1500.00")
    sc = rep.sales_collections()
    assert sc["net_sales"].sum() == Decimal("2900.00") and sc["receipts"].sum() == Decimal("2950.00")
    cf = rep.cash_flow_summary("01-10-2026", "31-10-2026")
    assert cf["sales_receipts"].sum() == Decimal("2950.00") and cf["supplier_payments"].sum() == Decimal("500.00")
    d = svc.dashboard_summary(end="30-11-2026")
    assert d["supplier_payments"] == Decimal("500.00") and d["bills_entered"] == Decimal("2000.00")


def test_cash_flow_opening_and_closing(svc, rep, vendor):
    svc.config["opening_cash"] = "1000"
    svc.config["opening_cash_date"] = "01-10-2026"
    day(svc, business_date="01-10-2026", gross_sales=500, discounts=0, cash=500, upi=0, card=0, tips=0)
    day(svc, business_date="02-10-2026", gross_sales=300, discounts=0, cash=300, upi=0, card=0, tips=0)
    out = rep.cash_flow_summary("01-10-2026", "02-10-2026")
    assert out["cash_closing"].tolist() == [Decimal("1500.00"), Decimal("1800.00")]
    later = rep.cash_flow_summary("02-10-2026", "02-10-2026")
    assert later["cash_closing"].tolist() == [Decimal("1800.00")]


def test_report_export_is_formula_safe(svc, rep, tmp_path):
    from openpyxl import load_workbook
    svc.add_vendor("=cmd|' /C calc'!A0", "Other")
    b = svc.add_bill(svc.db.read("vendors").iloc[0]["vendor_id"], "X", "01-10-2026", "02-10-2026", "Other", 10)
    out = rep.export(rep.vendor_outstanding("01-11-2026"), tmp_path / "r.xlsx", "Outstanding")
    ws = load_workbook(out).active
    assert all(c.data_type != "f" for row in ws.iter_rows() for c in row)
