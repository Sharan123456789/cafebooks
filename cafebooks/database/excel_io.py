"""Low-level, injection-safe xlsx reading and writing (openpyxl)."""
import re
from datetime import date, datetime
from decimal import Decimal
import math

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
MAX_CELL_CHARS = 32000


def clean_text(value):
    """Strip XML-illegal control characters and cap length."""
    return _ILLEGAL.sub("", str(value))[:MAX_CELL_CHARS]


def _text_cell(ws, value):
    cell = WriteOnlyCell(ws, value=value)
    # Force a *string* cell so Excel never evaluates user text such as "=HYPERLINK(...)".
    cell.data_type = "s"
    return cell


def write_xlsx(path, sheet_title, columns, rows, money_columns=(), widths=True):
    """Write a single-sheet workbook; every string is stored as literal text."""
    write_xlsx_multi(path, [(sheet_title, columns, rows, money_columns)], widths=widths)


def write_xlsx_multi(path, sheets, widths=True):
    """sheets: iterable of (title, columns, rows, money_columns)."""
    wb = Workbook(write_only=True)
    for title, columns, rows, money_columns in sheets:
        ws = wb.create_sheet(title=str(title)[:31])
        ws.freeze_panes = "A2"
        if widths:
            for idx, name in enumerate(columns, start=1):
                ws.column_dimensions[_col_letter(idx)].width = min(max(len(name) + 4, 12), 40)
        header = []
        for name in columns:
            c = WriteOnlyCell(ws, value=name)
            c.data_type = "s"
            c.font = Font(bold=True)
            c.fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
            header.append(c)
        ws.append(header)
        money_set = set(money_columns)
        money_idx = {i for i, name in enumerate(columns) if name in money_set}
        for row in rows:
            ws.append([_convert(ws, v, i in money_idx) for i, v in enumerate(row)])
    wb.save(path)


def _col_letter(idx):
    letters = ""
    while idx:
        idx, rem = divmod(idx - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _convert(ws, v, is_money):
    if v is None:
        return None
    if isinstance(v, float) and math.isnan(v):
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, Decimal) or (is_money and isinstance(v, (str, int, float))):
        if isinstance(v, str) and not v.strip():
            return None
        try:
            dv = v if isinstance(v, Decimal) else Decimal(str(v))
            cell = WriteOnlyCell(ws, value=float(dv))
            cell.number_format = "0.00"
            return cell
        except Exception:
            return _text_cell(ws, clean_text(v))
    if isinstance(v, (datetime, date)):
        return clean_text(v.strftime("%Y-%m-%d"))
    if isinstance(v, (int, float)):
        return v
    text = clean_text(v)
    if text[:1] in ("=", "+", "-", "@") or text[:1] == "\t":
        return _text_cell(ws, text)
    return text


def read_xlsx(path):
    """Return (header, rows) using cached values only (formulas are never evaluated)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        try:
            header = list(next(it))
        except StopIteration:
            return [], []
        keep = [i for i, h in enumerate(header) if h not in (None, "")]
        names = [str(header[i]).strip() for i in keep]
        rows = []
        for r in it:
            r = list(r)
            vals = [r[i] if i < len(r) else None for i in keep]
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in vals):
                continue
            rows.append(vals)
        return names, rows
    finally:
        wb.close()


def find_formula_cells(path, limit=10):
    """Return coordinates of real formula cells (used to reject untrusted imports)."""
    found = []
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if getattr(cell, "data_type", None) == "f":
                        found.append(f"{ws.title}!{cell.coordinate}")
                        if len(found) >= limit:
                            return found
    finally:
        wb.close()
    return found


def quick_verify(path, columns, nrows):
    """Fast post-write verification: zip CRCs, header row, and data-row count (streamed at C speed)."""
    import zipfile
    with zipfile.ZipFile(path) as z:
        bad = z.testzip()
        if bad:
            return f"corrupt archive member {bad}"
        sheets = [n for n in z.namelist() if n.startswith("xl/worksheets/sheet")]
        if len(sheets) != 1:
            return "unexpected sheet count"
        count = 0
        with z.open(sheets[0]) as fh:
            tail = b""
            while True:
                chunk = fh.read(4 * 1024 * 1024)
                if not chunk:
                    break
                buf = tail + chunk
                count += buf.count(b"<row ")
                tail = buf[-4:]
    if count != nrows + 1:
        return f"expected {nrows + 1} rows, found {count}"
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        first = next(wb.worksheets[0].iter_rows(values_only=True), None)
    finally:
        wb.close()
    if not first or [c for c in first if c is not None] != list(columns):
        return "header mismatch"
    return ""
