"""Init, vendors, bills, payments, aging, precision, security."""
import io
from decimal import Decimal
import pytest
from openpyxl import load_workbook
from cafebooks.database.schemas import SCHEMAS
from cafebooks.services.accounting_service import DuplicateError, DuplicateWarning
from cafebooks.utils.currency_utils import money


def test_first_run_creates_all_tables(db):
    for t in SCHEMAS:
        assert list(db.read(t).columns) == SCHEMAS[t], t
    assert db.get_setting("schema_version") == "4"


def test_reinitialization_never_erases_data(make_db):
    from cafebooks.services.accounting_service import AccountingService
    db = make_db()
    s = AccountingService(db)
    vid = s.add_vendor("Persist Me", "Other")
    db2 = make_db()                                      # simulated restart / second server
    assert vid in db2.read("vendors")["vendor_id"].tolist()


def test_vendor_create_deactivate_and_delete_rules(svc):
    v = svc.add_vendor("A Vendor", "Other", phone="9876543210", gstin="29ABCDE1234F1Z5")
    with pytest.raises(DuplicateWarning):
        svc.add_vendor("a  vendor", "Other")
    svc.add_vendor("A Vendor", "Other", allow_duplicate=True)         # legitimate similar vendor
    b = svc.add_bill(v, "X1", "01-02-2026", "10-02-2026", "Other", 100)
    with pytest.raises(ValueError, match="cannot be permanently deleted"):
        svc.delete_vendor(v)
    svc.deactivate_vendor(v)
    assert not any(vid == v for vid, _ in svc.vendor_options())
    with pytest.raises(ValueError, match="inactive"):
        svc.add_bill(v, "X2", "01-02-2026", "10-02-2026", "Other", 100)
    with pytest.raises(ValueError):
        svc.add_vendor("Bad", "Other", gstin="123")


def test_duplicate_invoice_detection_normalised(svc, vendor):
    svc.add_bill(vendor, "INV-001", "01-03-2026", "31-03-2026", "Other", 100)
    with pytest.raises(DuplicateError):
        svc.add_bill(vendor, " inv 001 ", "02-03-2026", "31-03-2026", "Other", 50)
    other = svc.add_vendor("Other V", "Other")
    svc.add_bill(other, "INV-001", "01-03-2026", "31-03-2026", "Other", 100)   # different vendor is fine


def test_invoice_total_must_reconcile(svc, vendor):
    with pytest.raises(ValueError, match="does not match"):
        svc.add_bill(vendor, "T1", "01-03-2026", "31-03-2026", "Other", 100, 18, 2, invoice_total=121)
    bid = svc.add_bill(vendor, "T2", "01-03-2026", "31-03-2026", "Other", 100, 18, 2, invoice_total=120)
    assert svc.bill_balances()[0]["invoice_total"] == "120.00"


def test_spec_example_partial_allocation(svc, vendor):
    a = svc.add_bill(vendor, "A", "01-03-2026", "31-03-2026", "Other", 10000)
    b = svc.add_bill(vendor, "B", "01-03-2026", "31-03-2026", "Other", 5000)
    svc.record_payment(vendor, "05-03-2026", 7000, "UPI", [(a, 5000), (b, 2000)])
    bal = {r["bill_id"]: r for r in svc.bill_balances("10-03-2026")}
    assert bal[a]["outstanding_balance"] == "5000.00" and bal[a]["calculated_status"] == "Partially Paid"
    assert bal[b]["outstanding_balance"] == "3000.00"
    assert svc.vendor_outstanding_total(vendor, "10-03-2026") == Decimal("8000.00")


def test_full_payment_and_status_precedence(svc, vendor):
    b = svc.add_bill(vendor, "F1", "01-01-2026", "10-01-2026", "Other", 500)
    assert svc.bill_balances("20-01-2026")[0]["calculated_status"] == "Overdue"
    svc.record_payment(vendor, "05-01-2026", 500, "Cash", [(b, 500)])
    assert svc.bill_balances("20-01-2026")[0]["calculated_status"] == "Paid"   # paid beats overdue


def test_rejects_overpayment_and_overallocation(svc, vendor):
    b = svc.add_bill(vendor, "O1", "01-03-2026", "31-03-2026", "Other", 1000)
    with pytest.raises(ValueError, match="exceed the payment"):
        svc.record_payment(vendor, "02-03-2026", 100, "Cash", [(b, 150)])
    with pytest.raises(ValueError, match="outstanding"):
        svc.record_payment(vendor, "02-03-2026", 2000, "Cash", [(b, 1500)])
    with pytest.raises(ValueError):
        svc.record_payment(vendor, "02-03-2026", 0, "Cash")
    with pytest.raises(ValueError):
        svc.record_payment(vendor, "02-03-2026", -5, "Cash")
    svc.record_payment(vendor, "02-03-2026", 700, "Cash", [(b, 700)])
    with pytest.raises(ValueError, match="outstanding"):
        svc.record_payment(vendor, "03-03-2026", 400, "Cash", [(b, 400)])
    assert len(svc.db.read("payments")) == 1          # failed attempts wrote nothing


def test_unallocated_payment_and_later_allocation(svc, vendor):
    b = svc.add_bill(vendor, "U1", "01-03-2026", "31-03-2026", "Other", 1000)
    p = svc.record_payment(vendor, "02-03-2026", 600, "Cash")
    assert svc.unallocated_total("05-03-2026") == Decimal("600.00")
    assert svc.bill_balances("05-03-2026")[0]["outstanding_balance"] == "1000.00"   # not settled
    svc.allocate_payment(p, [(b, 250)], "06-03-2026")
    assert svc.unallocated_total("07-03-2026") == Decimal("350.00")
    assert svc.bill_balances("07-03-2026")[0]["outstanding_balance"] == "750.00"
    with pytest.raises(ValueError):
        svc.allocate_payment(p, [(b, 400)], "06-03-2026")


def test_reversal_keeps_history_and_is_audited(svc, vendor):
    b = svc.add_bill(vendor, "R1", "01-03-2026", "31-03-2026", "Other", 1000)
    p = svc.record_payment(vendor, "02-03-2026", 1000, "Cash", [(b, 1000)])
    svc.reverse_payment(p, "wrong vendor", "10-03-2026")
    alloc = svc.db.read("payment_allocations")
    assert list(alloc["entry_type"]) == ["ALLOCATE", "REVERSE"]            # nothing flipped or deleted
    assert alloc.iloc[0]["status"] == "Active"
    assert svc.bill_balances("05-03-2026")[0]["outstanding_balance"] == "0.00"     # history intact
    assert svc.bill_balances("15-03-2026")[0]["outstanding_balance"] == "1000.00"
    with pytest.raises(ValueError):
        svc.reverse_payment(p, "again")
    audit = svc.db.read("audit_log")
    assert "reverse" in audit["action"].tolist()
    row = audit[audit["action"] == "reverse"].iloc[0]
    assert row["actor"] == "tester" and "Reversed" in row["after"]


def test_cancel_blocked_while_paid_then_allowed_after_reversal(svc, vendor):
    b = svc.add_bill(vendor, "C1", "01-03-2026", "31-03-2026", "Other", 1000)
    p = svc.record_payment(vendor, "02-03-2026", 400, "Cash", [(b, 400)])
    with pytest.raises(ValueError, match="Reverse"):
        svc.cancel_bill(b, "oops")
    svc.reverse_payment(p, "undo")
    svc.cancel_bill(b, "oops")
    with pytest.raises(ValueError, match="cancelled"):
        svc.record_payment(vendor, "05-03-2026", 100, "Cash", [(b, 100)])


def test_historical_aging_uses_as_of_not_today(svc, rep, vendor):
    b = svc.add_bill(vendor, "H1", "01-01-2026", "31-01-2026", "Other", 1000)
    svc.record_payment(vendor, "01-03-2026", 1000, "Cash", [(b, 1000)])
    aging = rep.vendor_aging("15-02-2026")
    assert len(aging) == 1 and aging.iloc[0]["days_overdue"] == 15 and aging.iloc[0]["aging_bucket"] == "1-30 days"
    assert aging.iloc[0]["outstanding_balance"] == Decimal("1000.00")
    assert rep.vendor_aging("02-03-2026").empty
    assert rep.vendor_aging("15-01-2026").iloc[0]["aging_bucket"] == "Current"


def test_aging_buckets_and_future_bills_excluded(svc, rep, vendor):
    for n, due in enumerate(["01-01-2026", "01-02-2026", "01-03-2026", "01-04-2026"]):
        svc.add_bill(vendor, f"B{n}", "01-12-2025", due, "Other", 100)
    svc.add_bill(vendor, "FUT", "01-06-2026", "30-06-2026", "Other", 999)   # did not exist yet
    s = rep.vendor_aging_summary("15-04-2026")
    total = s[s.vendor_name == "TOTAL"].iloc[0]
    assert total["total"] == Decimal("400.00")
    assert total[">90 days"] == Decimal("100.00") and total["1-30 days"] == Decimal("100.00")


def test_paise_precision(svc, vendor):
    assert money("0.1") + money("0.2") == Decimal("0.30")
    with pytest.raises(ValueError):
        money("10.005")
    b = svc.add_bill(vendor, "P1", "01-03-2026", "31-03-2026", "Other", "100.10", "0.20", "0.01")
    svc.record_payment(vendor, "02-03-2026", "50.07", "Cash", [(b, "50.07")])
    assert svc.bill_balances("03-03-2026")[0]["outstanding_balance"] == "50.24"


def test_update_bill_whitelist_and_locking(svc, vendor):
    b = svc.add_bill(vendor, "W1", "01-03-2026", "31-03-2026", "Other", 100)
    with pytest.raises(ValueError):
        svc.update_bill(b, status="Paid")
    with pytest.raises(ValueError):
        svc.update_bill(b, invoice_total="1")
    svc.record_payment(vendor, "02-03-2026", 50, "Cash", [(b, 50)])
    with pytest.raises(ValueError, match="allocated"):
        svc.update_bill(b, subtotal=200)
    svc.update_bill(b, description="fixed", reason="typo")
    audit = svc.db.read("audit_log")
    assert audit[audit["action"] == "update"].iloc[-1]["before"] != ""


def test_formula_text_is_stored_as_text(svc, db):
    svc.add_vendor('=HYPERLINK("http://x","a")', "Other", notes="=1+1")
    wb = load_workbook(io.BytesIO(db.export_database_snapshot()))
    assert all(c.data_type != "f" for ws in wb.worksheets for row in ws.iter_rows() for c in row)
    assert db.read("vendors").iloc[0]["name"].startswith("=HYPERLINK")


def test_idempotent_retry_does_not_duplicate(svc, vendor):
    b = svc.add_bill(vendor, "ID1", "01-03-2026", "31-03-2026", "Other", 100, request_id="req-1")
    again = svc.add_bill(vendor, "ID1b", "01-03-2026", "31-03-2026", "Other", 100, request_id="req-1")
    assert b == again and len(svc.db.read("bills")) == 1
    p1 = svc.record_payment(vendor, "02-03-2026", 40, "Cash", [(b, 40)], request_id="pay-1")
    p2 = svc.record_payment(vendor, "02-03-2026", 40, "Cash", [(b, 40)], request_id="pay-1")
    assert p1 == p2 and len(svc.db.read("payments")) == 1 and len(svc.db.read("payment_allocations")) == 1


def test_cache_matches_database_after_operations(svc, vendor, make_db):
    b = svc.add_bill(vendor, "CC1", "01-03-2026", "31-03-2026", "Other", 100)
    svc.record_payment(vendor, "02-03-2026", 40, "UPI", [(b, 40)])
    fresh = make_db()
    for t in SCHEMAS:
        assert fresh.read(t).equals(svc.db.read(t)), t
