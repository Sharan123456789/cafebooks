"""End-to-end tests of the web API (the same app Vercel and the desktop app run)."""
import io
import json
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from cafebooks.web import server

DEMO = Path(__file__).resolve().parents[1] / "demo_data" / "CafeBooks_Demo_Data.xlsx"


@pytest.fixture
def client(db_url, tmp_path, monkeypatch):
    for key in ("DATABASE_URL", "POSTGRES_URL", "VERCEL", "CRON_SECRET", "CAFEBOOKS_MODE"):
        monkeypatch.delenv(key, raising=False)
    monkeypatch.setenv("CAFEBOOKS_DATABASE_URL", db_url)
    monkeypatch.setenv("CAFEBOOKS_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("CAFEBOOKS_BACKUP_STORE", "database")
    server.reset_runtime()
    with TestClient(server.create_app(), headers={"X-CafeBooks-User": "Priya%20(counter)",
                                                  "X-Requested-With": "CafeBooks"}) as c:
        yield c
    server.reset_runtime()


def test_forged_cross_site_requests_are_refused(client, monkeypatch, tmp_path):
    forged = client.post("/api/vendors", json={"name": "X", "category": "Other"}, headers={"X-Requested-With": ""})
    assert forged.status_code == 403
    upload = client.post("/api/data/import-file/apply", headers={"X-Requested-With": ""},
                         files={"file": ("d.xlsx", DEMO.read_bytes(), "application/octet-stream")})
    assert upload.status_code == 403 and ok(client.get("/api/vendors")) == []
    monkeypatch.setenv("CAFEBOOKS_MODE", "desktop")              # desktop: only its own local address
    server.reset_runtime()
    assert client.get("/api/vendors", headers={"Host": "evil.example.com"}).status_code == 403
    assert client.get("/api/vendors", headers={"Host": "127.0.0.1:8765"}).status_code == 200
    exe = tmp_path / "tool.exe"
    exe.write_bytes(b"MZ")
    v = ok(client.post("/api/vendors", json={"name": "V", "category": "Other"}, headers={"Host": "localhost"}))["vendor_id"]
    b = ok(client.post("/api/bills", headers={"Host": "localhost"}, json={
        "vendor_id": v, "invoice_number": "1", "invoice_date": "2026-01-01", "due_date": "2026-01-02",
        "category": "Other", "subtotal": "1", "attachment_path": str(exe)}))["bill_id"]
    refused = client.post(f"/api/bills/{b}/open-attachment", headers={"Host": "localhost"})
    assert refused.status_code == 400 and "Only documents" in refused.json()["error"]


def ok(resp, status=200):
    assert resp.status_code == status, resp.text
    body = resp.json()
    assert body["ok"] is True, body
    return body.get("data")


def test_ui_and_bootstrap(client):
    page = client.get("/")
    assert page.status_code == 200 and "CafeBooks" in page.text and 'name="robots"' in page.text
    assert client.get("/js/app.js").status_code == 200
    boot = ok(client.get("/api/bootstrap"))
    assert boot["config"]["cafe_name"] == "My Cafe" and boot["mode"] == "web"
    assert len(boot["reports"]) == 17 and "UPI" in boot["payment_methods"]
    assert client.get("/api/health").json()["ok"] is True


def test_full_day_of_bookkeeping(client):
    v = ok(client.post("/api/vendors", json={"name": "Fresh Milk Co", "category": "Milk and Dairy", "payment_terms_days": 10}))["vendor_id"]
    dup = client.post("/api/vendors", json={"name": "fresh  milk co", "category": "Other"})
    assert dup.status_code == 409 and dup.json()["confirm"] == "allow_duplicate"
    ok(client.post("/api/vendors", json={"name": "fresh milk co", "category": "Other", "allow_duplicate": True}))

    b1 = ok(client.post("/api/bills", json={"vendor_id": v, "invoice_number": "INV-1", "invoice_date": "2026-10-01",
                                            "due_date": "2026-10-11", "category": "Milk and Dairy", "subtotal": "1000",
                                            "tax_amount": "50.50", "request_id": "r-bill-1"}))["bill_id"]
    again = ok(client.post("/api/bills", json={"vendor_id": v, "invoice_number": "INV-1x", "invoice_date": "2026-10-01",
                                               "due_date": "2026-10-11", "category": "Other", "subtotal": "1",
                                               "request_id": "r-bill-1"}))["bill_id"]
    assert again == b1                                                    # retried submit, no duplicate
    clash = client.post("/api/bills", json={"vendor_id": v, "invoice_number": "inv 1", "invoice_date": "2026-10-02",
                                            "due_date": "2026-10-12", "category": "Other", "subtotal": "5"})
    assert clash.status_code == 409 and "confirm" not in clash.json()
    b2 = ok(client.post("/api/bills", json={"vendor_id": v, "invoice_number": "INV-2", "invoice_date": "2026-10-03",
                                            "due_date": "2026-10-13", "category": "Other", "subtotal": "400",
                                            "attachment_path": "https://drive.example.com/inv2.pdf"}))["bill_id"]
    open_bills = ok(client.get(f"/api/bills/open?vendor_id={v}&as_of=2026-10-05"))
    assert [b["bill_id"] for b in open_bills] == [b1, b2]

    pay = ok(client.post("/api/payments", json={"vendor_id": v, "payment_date": "2026-10-05", "amount": "1200",
                                                "method": "UPI", "allocations": [{"bill_id": b1, "amount": "1050.50"},
                                                                                 {"bill_id": b2, "amount": "100"}]}))["payment_id"]
    bills = {b["bill_id"]: b for b in ok(client.get("/api/bills?as_of=2026-10-06"))}
    assert bills[b1]["calculated_status"] == "Paid" and bills[b2]["outstanding_balance"] == "300.00"
    p = ok(client.get("/api/payments"))[0]
    assert p["unallocated"] == "49.50"
    ok(client.post(f"/api/payments/{pay}/allocate", json={"allocations": [[b2, "49.50"]], "allocation_date": "2026-10-06"}))
    allocs = ok(client.get(f"/api/payments/{pay}/allocations"))
    assert len(allocs) == 3 and allocs[0]["invoice_number"] == "INV-1"
    over = client.post(f"/api/allocations/{allocs[0]['allocation_id']}/reverse", json={"reason": "x", "amount": "9999"})
    assert over.status_code == 400
    ok(client.post(f"/api/allocations/{allocs[1]['allocation_id']}/reverse", json={"reason": "wrong bill"}))
    bad_cancel = client.post(f"/api/bills/{b1}/cancel", json={"reason": "oops"})
    assert bad_cancel.status_code == 400 and "Reverse" in bad_cancel.json()["error"]
    ok(client.post(f"/api/payments/{pay}/reverse", json={"reason": "bank bounced", "reversal_date": "2026-10-07"}))
    ok(client.post(f"/api/bills/{b1}/cancel", json={"reason": "duplicate entry"}))
    ok(client.patch(f"/api/bills/{b2}", json={"description": "Paneer", "reason": "typo"}))

    s = ok(client.post("/api/sales", json={"business_date": "2026-10-05", "gross_sales": "5000", "discounts": "200",
                                           "cash_collected": "2000", "upi_collected": "2850", "tips_collected": "50"}))["sales_id"]
    dup_day = client.post("/api/sales", json={"business_date": "2026-10-05", "gross_sales": "1"})
    assert dup_day.status_code == 409 and dup_day.json()["confirm"] == "allow_duplicate"
    ok(client.patch(f"/api/sales/{s}", json={"actual_cash": "1990", "reason": "counted"}))
    sale = ok(client.get("/api/sales"))[0]
    assert sale["net_sales"] == "4800.00" and sale["cash_difference"] == "-10.00"
    ok(client.post(f"/api/sales/{s}/adjustments", json={"adjustment_type": "REFUND", "amount": "100", "adjustment_date": "2026-10-06"}))
    adj = ok(client.get("/api/adjustments"))[0]
    tip = ok(client.get("/api/tips"))[0]
    assert tip["collected"] == "50.00"
    ok(client.patch(f"/api/tips/{tip['entry_id']}", json={"distributed": "50", "retained": "0", "distribution_date": "2026-10-06"}))
    ok(client.post(f"/api/adjustments/{adj['adjustment_id']}/void", json={"reason": "entered twice"}))

    dash = ok(client.get("/api/dashboard?start=2026-10-01&end=2026-10-31"))
    assert dash["summary"]["net_sales"] == "4800.00" and dash["summary"]["collections_by_method"]["UPI"] == "2850.00"
    assert dash["integrity"]["errors"] == 0

    activity = ok(client.get("/api/activity?limit=5"))
    assert activity[0]["actor"] == "Priya (counter)"


def test_every_report_runs_and_exports(client):
    v = ok(client.post("/api/vendors", json={"name": "V", "category": "Other"}))["vendor_id"]
    ok(client.post("/api/bills", json={"vendor_id": v, "invoice_number": "A", "invoice_date": "2026-09-01",
                                       "due_date": "2026-09-05", "category": "Other", "subtotal": "100"}))
    ok(client.post("/api/sales", json={"business_date": "2026-09-02", "gross_sales": "300", "cash_collected": "300"}))
    for r in ok(client.get("/api/bootstrap"))["reports"]:
        data = ok(client.get("/api/reports/run", params={"name": r["name"], "start": "2026-08-01", "end": "2026-10-01"}))
        assert data["columns"], r["name"]
    out = ok(client.get("/api/reports/run", params={"name": "Aging - bill wise", "end": "2026-10-01"}))
    assert out["rows"][0][out["columns"].index("aging_bucket")] == "1-30 days"
    x = client.get("/api/reports/export", params={"name": "Expenses by category", "start": "2026-08-01", "end": "2026-10-01"})
    assert x.status_code == 200 and "attachment" in x.headers["content-disposition"]
    assert load_workbook(io.BytesIO(x.content)).active["A1"].value == "category"
    assert client.get("/api/reports/run", params={"name": "Nope"}).status_code == 400


def test_settings_are_validated_shared_and_audited(client):
    cfg = ok(client.put("/api/settings", json={"cafe_name": "Sunrise Bakery", "vendor_categories": ["Flour", "Butter"],
                                               "opening_cash": "1500"}))
    assert cfg["cafe_name"] == "Sunrise Bakery" and cfg["opening_cash"] == "1500.00"
    assert ok(client.get("/api/settings"))["vendor_categories"] == ["Flour", "Butter"]
    bad = client.put("/api/settings", json={"date_display": "MM/DD/YYYY"})
    assert bad.status_code == 400
    assert client.put("/api/settings", json={"timezone": "Nowhere/Land"}).status_code == 400
    assert ok(client.get("/api/activity"))[0]["entity_type"] == "settings"


def test_backups_download_upload_and_restore(client):
    ok(client.post("/api/vendors", json={"name": "Before", "category": "Other"}))
    name = ok(client.post("/api/backups"))["name"]
    listing = ok(client.get("/api/backups"))
    assert listing["store"] == "database" and listing["items"][0]["name"] == name
    zipped = client.get(f"/api/backups/{name}/download")
    assert zipped.status_code == 200 and zipfile.ZipFile(io.BytesIO(zipped.content)).testzip() is None
    ok(client.post("/api/vendors", json={"name": "After", "category": "Other"}))
    preview = ok(client.post(f"/api/backups/{name}/preview"))
    assert preview["tables"]["vendors"] == {"backup": 1, "current": 2} and preview["errors"] == []
    up = ok(client.post("/api/restore-upload/preview", files={"file": ("b.zip", zipped.content, "application/zip")}))
    assert up["tables"]["vendors"]["backup"] == 1
    safety = ok(client.post("/api/restore-upload/apply", files={"file": ("b.zip", zipped.content, "application/zip")}))["safety_backup"]
    assert [v["name"] for v in ok(client.get("/api/vendors"))] == ["Before"]
    assert safety in [b["name"] for b in ok(client.get("/api/backups"))["items"]]
    assert client.get("/api/backups/..%2F..%2Fsecret.zip/download").status_code in (400, 404)
    junk = client.post("/api/restore-upload/preview", files={"file": ("x.zip", b"not a zip", "application/zip")})
    assert junk.status_code == 400


def test_excel_tools(client):
    ok(client.post("/api/vendors", json={"name": "Packer", "category": "Other"}))
    snap = client.get("/api/data/snapshot")
    assert snap.status_code == 200 and "vendors" in load_workbook(io.BytesIO(snap.content)).sheetnames

    pack = client.get("/api/data/edit-pack")
    assert pack.status_code == 200
    src = zipfile.ZipFile(io.BytesIO(pack.content))
    edited = io.BytesIO()
    with zipfile.ZipFile(edited, "w") as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.endswith("EDIT_THESE_FILES/vendors.xlsx"):
                wb = load_workbook(io.BytesIO(data))
                wb.active["B2"] = "Packer & Sons"
                buf = io.BytesIO()
                wb.save(buf)
                data = buf.getvalue()
            out.writestr(item.filename, data)
    files = {"file": ("pack.zip", edited.getvalue(), "application/zip")}
    plan = ok(client.post("/api/data/edit-pack/preview", files=files))
    assert plan["errors"] == [] and plan["counts"]["vendors"]["modified"] == 1
    res = ok(client.post("/api/data/edit-pack/apply", files=files, data={"resolutions": json.dumps([])}))
    assert res["applied"] and ok(client.get("/api/vendors"))[0]["name"] == "Packer & Sons"

    demo = client.get("/api/data/demo-workbook")
    assert demo.status_code == 200
    demo_file = {"file": ("CafeBooks_Demo_Data.xlsx", demo.content, "application/vnd.ms-excel")}
    preview = ok(client.post("/api/data/import-file/preview", files=demo_file))
    assert preview["errors"] == [] and preview["tables"]["vendors"]["incoming"] == 2
    ok(client.post("/api/data/import-file/apply", files=demo_file))
    assert len(ok(client.get("/api/vendors"))) == 2
    assert ok(client.get("/api/integrity")) == []

    assert client.post("/api/data/reset", json={"confirm": "yes"}).status_code == 400
    backup = ok(client.post("/api/data/reset", json={"confirm": "RESET"}))["backup"]
    assert ok(client.get("/api/vendors")) == [] and backup


def test_daily_backup_cron(client, monkeypatch):
    first = ok(client.get("/api/cron/daily-backup"))
    assert first["backup"].startswith("daily_")
    assert ok(client.get("/api/cron/daily-backup"))["backup"] is None
    monkeypatch.setenv("CRON_SECRET", "s3cret")
    server.reset_runtime()
    assert client.get("/api/cron/daily-backup").status_code == 401
    assert client.get("/api/cron/daily-backup", headers={"Authorization": "Bearer s3cret"}).status_code == 200


def test_vercel_without_database_shows_setup_message(monkeypatch, tmp_path):
    for key in ("CAFEBOOKS_DATABASE_URL", "DATABASE_URL", "POSTGRES_URL"):
        monkeypatch.delenv(key, raising=False)
    monkeypatch.setenv("VERCEL", "1")
    server.reset_runtime()
    with TestClient(server.create_app()) as c:
        r = c.get("/api/bootstrap")
        assert r.status_code == 503 and r.json()["setup"] is True and "DATABASE_URL" in r.json()["error"]
        assert c.get("/api/health").status_code == 503
        assert c.get("/").status_code == 200                     # the UI still loads and explains the problem
    server.reset_runtime()
